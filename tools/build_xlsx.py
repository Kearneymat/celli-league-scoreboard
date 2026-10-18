import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import FormulaRule
WEEK = sys.argv[1] if len(sys.argv) > 1 else '5'
STATUS = sys.argv[2] if len(sys.argv) > 2 else 'all games final'
OUT = sys.argv[3]
d = json.load(open('page_data.json'))
teams = sorted(d['teams'], key=lambda t: -t['total'])
wb = Workbook()
F = Font(name='Arial', size=10); FB = Font(name='Arial', size=10, bold=True); FH = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FT = Font(name='Arial', size=14, bold=True); FS = Font(name='Arial', size=9, italic=True, color='5D6A62')
HF = PatternFill('solid', fgColor='1D6B45'); GOLD = PatternFill('solid', fgColor='FBEFD2'); CUT = PatternFill('solid', fgColor='FDE7E2')
thin = Border(bottom=Side(style='thin', color='D3D9D0'))
def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row, i, c); x.font = FH; x.fill = HF; x.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[CL(i)].width = w
    ws.freeze_panes = ws.cell(row + 1, 1)
def title(ws, t, sub):
    ws['A1'] = t; ws['A1'].font = FT; ws['A2'] = sub; ws['A2'].font = FS
L = wb.active; L.title = 'Lineups'
title(L, f'Week {WEEK} Lineups', f'Every team’s 15 players and how each one scored · Week {WEEK} · {STATUS} · unofficial until the commissioner posts results')
header(L, 4, ['Owner', 'Pos', 'Player', 'Game', 'Pts', 'How it scored'], [18, 6, 24, 20, 7, 80])
r = 5
for t in teams:
    for p in t['lineup']:
        bd = ' · '.join(f"{a} ({'+' if v > 0 else ''}{v})" for a, v in p['bd']) or 'No scoring plays'
        for i, v in enumerate([t['owner'], p['pos'], p['name'], p['game'] or p['team'], p['pts'], bd], 1):
            c = L.cell(r, i, v); c.font = F; c.border = thin
        r += 1
LAST = r - 1
L.auto_filter.ref = f"A4:F{LAST}"
rngO, rngP, rngN, rngPos = (f"Lineups!${c}$5:${c}${LAST}" for c in 'AECB')
W = wb.create_sheet('Leaderboard', 0)
title(W, f'Salary Cap League · Week {WEEK} Leaderboard', f'Week {WEEK} · {STATUS} · unofficial until the commissioner posts results')
header(W, 4, ['Rank', 'Owner', f'Week {WEEK}', f'Before Week {WEEK}', 'Season total', 'Season rank', 'Showdown'], [7, 20, 10, 14, 13, 11, 11])
a, b = 5, 4 + len(teams)
for i, t in enumerate(teams):
    rr = 5 + i
    W.cell(rr, 2, t['owner']); W.cell(rr, 3, f'=SUMIFS({rngP},{rngO},B{rr})'); W.cell(rr, 4, t['prev'])
    W.cell(rr, 5, f'=D{rr}+C{rr}'); W.cell(rr, 1, f'=RANK(C{rr},$C${a}:$C${b})'); W.cell(rr, 6, f'=RANK(E{rr},$E${a}:$E${b})')
    W.cell(rr, 7, 'Alive' if t['alive'] else 'Out')
    for c in range(1, 8):
        W.cell(rr, c).font = FB if c in (2, 3) else F; W.cell(rr, c).border = thin
        if i == 0: W.cell(rr, c).fill = GOLD
W.cell(b + 2, 1, f'"Before Week {WEEK}" comes from the commissioner’s sheet. Showdown status = owners not in red on that sheet.').font = FS
W.auto_filter.ref = f"A4:G{b}"
S = wb.create_sheet('Showdown', 1)
alive = [t for t in teams if t['alive']]
title(S, f'Showdown · Week {WEEK}', 'Teams still alive, ranked by this week only. The bottom number in C3 are cut; change C3 if the commissioner sets a different number.')
S['B3'] = 'Cut this week:'; S['B3'].font = FB; S['C3'] = 6; S['C3'].font = Font(name='Arial', size=10, bold=True, color='0000FF')
header(S, 5, ['Rank', 'Owner', f'Week {WEEK}', 'Status'], [7, 20, 10, 14])
a2, b2 = 6, 5 + len(alive)
for i, t in enumerate(alive):
    rr = 6 + i
    S.cell(rr, 2, t['owner']); S.cell(rr, 3, f'=SUMIFS({rngP},{rngO},B{rr})'); S.cell(rr, 1, f'=RANK(C{rr},$C${a2}:$C${b2})')
    S.cell(rr, 4, f'=IF(A{rr}>COUNT($C${a2}:$C${b2})-$C$3,"Cut zone","Safe")')
    for c in range(1, 5): S.cell(rr, c).font = F; S.cell(rr, c).border = thin
S.conditional_formatting.add(f"A{a2}:D{b2}", FormulaRule(formula=[f'$D{a2}="Cut zone"'], fill=CUT))
P = wb.create_sheet('Players')
title(P, f'Players · Week {WEEK}', 'Every rostered player, his points, and how many of the 50 teams used him')
header(P, 4, ['Pos', 'Player', 'Game', 'Pts', 'Teams', 'How it scored'], [6, 24, 20, 7, 7, 80])
seen = {}
for t in teams:
    for p in t['lineup']: seen.setdefault((p['pos'], p['name']), p)
for i, p in enumerate(sorted(seen.values(), key=lambda p: -p['pts'])):
    rr = 5 + i
    bd = ' · '.join(f"{a} ({'+' if v > 0 else ''}{v})" for a, v in p['bd']) or 'No scoring plays'
    for c, v in enumerate([p['pos'], p['name'], p['game'] or p['team'], p['pts'], f'=COUNTIFS({rngN},B{rr},{rngPos},A{rr})', bd], 1):
        x = P.cell(rr, c, v); x.font = F; x.border = thin
P.auto_filter.ref = f"A4:F{4 + len(seen)}"
R = wb.create_sheet('Scoring rules')
title(R, 'Scoring rules used', 'From the league scoring sheet; lost fumbles only, −3')
rules = [('Touchdowns (all)', '0–9 yds 6 · 10–19 7 · 20–29 8 · 30–39 9 · 40–49 10 · 50+ 12'), ('QB passing yards', '200+ = 5, +1 per 10 yds over 200'),
 ('QB completions', '20 = 5, +1 per completion over 20'), ('Interceptions', '−5'), ('Lost fumbles', '−3'), ('Rushing yards (all)', '50 = 5, +1 per 10 yds over'),
 ('Carries (QB, RB)', '15 = 5, +1 per carry over 15'), ('Receptions (RB, WR)', '5 = 5, +1 per catch over 5'), ('Receptions (TE)', '3 = 5, +2 per catch over 3'),
 ('Receiving yards (RB, WR)', '50 = 5, +1 per 10 yds over'), ('Receiving yards (TE)', '30 = 5, +2 per 10 yds over'), ('2-pt conversion', '2'),
 ('Field goals', '0–29 3 · 30–39 6 · 40–49 9 · 50–59 12 · 60+ 15'), ('Missed field goals', '0–29 −9 · 30–39 −6 · 40–49 −3 · 50–59 −1 · 60+ 0'),
 ('Extra points', 'made +1 · missed −9'), ('Defense points allowed', '0: 20 · 1–6: 15 · 7–10: 10 · 11–14: 5 · 15–24: 0 · 25–30: −5 · 31–35: −10 · 36–44: −15 · 45+: −20'),
 ('Defense other', 'safety 5 · sack 2 · INT 2 · fumble recovery 2 · win +5 · loss −5 · return/defensive TD by length')]
header(R, 4, ['Category', 'Points'], [26, 90]); R.freeze_panes = None
for i, (k, v) in enumerate(rules): R.cell(5 + i, 1, k).font = FB; R.cell(5 + i, 2, v).font = F
wb.save(OUT); print(OUT)
